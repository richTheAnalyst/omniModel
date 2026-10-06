"""PDF and DOCX export for proposal letters and emails."""

from __future__ import annotations

import io


def export_pdf(text: str) -> bytes:
    """Export text as a PDF using reportlab."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    width, height = letter

    y = height - 50
    for line in text.split("\n"):
        if y < 50:
            c.showPage()
            y = height - 50
        c.drawString(50, y, line)
        y -= 14

    c.save()
    buf.seek(0)
    return buf.getvalue()


def export_docx(text: str, title: str = "Proposal") -> bytes:
    """Export text as a DOCX using python-docx."""
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    for line in text.split("\n"):
        if line.strip():
            doc.add_paragraph(line)
        else:
            doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def export_txt(text: str) -> bytes:
    """Export text as plain text."""
    return text.encode("utf-8")
